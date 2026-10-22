# -*- coding: utf-8 -*-
"""Builds the client-facing "Behind the Assistant" brief as a .docx file."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

INK = RGBColor(0x17, 0x26, 0x2A)
MUTED = RGBColor(0x56, 0x69, 0x6C)
TEAL = RGBColor(0x0F, 0x6E, 0x66)
TEAL_DEEP = RGBColor(0x0A, 0x4B, 0x46)
GOLD = RGBColor(0xB8, 0x86, 0x3F)
TEAL_TINT = "E3F1EE"
GOLD_TINT = "F5ECDB"
LINE_GREY = "DCE6E4"

HEADING_FONT = "Georgia"
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

doc = Document()

# ---- base document setup -----------------------------------------------
section = doc.sections[0]
section.page_width = Inches(8.27)   # A4
section.page_height = Inches(11.69)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.top_margin = Inches(0.9)
section.bottom_margin = Inches(0.9)

normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(11)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.28


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color=LINE_GREY, size=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def add_para(text="", size=11, color=INK, bold=False, italic=False,
             font=BODY_FONT, space_before=0, space_after=8, align=None,
             indent=None, line_spacing=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if line_spacing:
        p.paragraph_format.line_spacing = line_spacing
    if align:
        p.alignment = align
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if text:
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
    return p


def add_mixed_para(parts, size=11, space_before=0, space_after=10,
                    line_spacing=1.3, indent=None):
    """parts: list of (text, {bold, italic, color, font, size})"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    for text, style in parts:
        run = p.add_run(text)
        run.font.size = Pt(style.get("size", size))
        run.font.color.rgb = style.get("color", INK)
        run.font.bold = style.get("bold", False)
        run.font.italic = style.get("italic", False)
        run.font.name = style.get("font", BODY_FONT)
    return p


def add_heading2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = HEADING_FONT
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = TEAL_DEEP
    return p


def add_eyebrow(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    run.font.name = MONO_FONT
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.color.rgb = TEAL
    run.font.all_caps = True
    add_letter_spacing(run, 16)
    return p


def add_letter_spacing(run, val_twentieths):
    rPr = run._element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(val_twentieths))
    rPr.append(spacing)


def add_hr():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), LINE_GREY)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_question_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(16)
    p.paragraph_format.left_indent = Inches(0.22)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), "0F6E66")
    pBdr.append(left)
    pPr.append(pBdr)
    run = p.add_run(text)
    run.font.name = HEADING_FONT
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = INK
    return p


def add_bullets(items):
    for label, desc in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(6)
        r1 = p.add_run(label + "  ")
        r1.bold = True
        r1.font.color.rgb = TEAL_DEEP
        r1.font.size = Pt(11)
        r1.font.name = BODY_FONT
        r2 = p.add_run(desc)
        r2.font.color.rgb = MUTED
        r2.font.size = Pt(11)
        r2.font.name = BODY_FONT


def add_flow(steps):
    table = doc.add_table(rows=1, cols=len(steps))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, step in enumerate(steps):
        cell = table.cell(0, i)
        set_cell_shading(cell, TEAL_TINT)
        set_cell_borders(cell)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        num_run = cell.paragraphs[0].add_run(f"{i+1:02d}")
        num_run.font.name = MONO_FONT
        num_run.font.size = Pt(9)
        num_run.font.bold = True
        num_run.font.color.rgb = TEAL
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(step)
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = INK
        r2.font.name = BODY_FONT
        for para in cell.paragraphs:
            for run in para.runs:
                pass
    add_para("", size=4, space_after=10)


def add_spec_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(1.5), Inches(1.9), Inches(2.6)]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr_cells[i], TEAL_TINT)
        set_cell_borders(hdr_cells[i])
        hdr_cells[i].width = widths[i] if i < len(widths) else Inches(2)
        run = hdr_cells[i].paragraphs[0].add_run(h.upper())
        run.font.name = MONO_FONT
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = TEAL_DEEP
        add_letter_spacing(run, 10)

    for r, row in enumerate(rows, start=1):
        cells = table.rows[r].cells
        for i, val in enumerate(row):
            cells[i].width = widths[i] if i < len(widths) else Inches(2)
            set_cell_borders(cells[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(val)
            run.font.size = Pt(9.5)
            run.font.name = BODY_FONT
            if i == 0:
                run.font.name = MONO_FONT
                run.font.size = Pt(9)
                run.font.color.rgb = TEAL_DEEP
                run.font.bold = True
            elif i == 1:
                run.font.bold = True
                run.font.color.rgb = INK
            else:
                run.font.color.rgb = MUTED
    add_para("", size=4, space_after=10)


def add_glossary(title, items):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, GOLD_TINT)
    set_cell_borders(cell, color="E4CE9B", size=4)
    cell.paragraphs[0].paragraph_format.space_after = Pt(6)
    t_run = cell.paragraphs[0].add_run(title.upper())
    t_run.font.name = MONO_FONT
    t_run.font.size = Pt(8.5)
    t_run.font.bold = True
    t_run.font.color.rgb = GOLD
    add_letter_spacing(t_run, 10)

    for term, definition in items:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(term + " — ")
        r1.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = INK
        r1.font.name = BODY_FONT
        r2 = p.add_run(definition)
        r2.font.size = Pt(10)
        r2.font.color.rgb = MUTED
        r2.font.name = BODY_FONT
    add_para("", size=4, space_after=14)


# ============================= COVER / HEADER =============================
add_eyebrow("Project brief")

title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(4)
title_p.paragraph_format.space_after = Pt(6)
title_run = title_p.add_run("Behind the Assistant")
title_run.font.name = HEADING_FONT
title_run.font.size = Pt(30)
title_run.font.bold = True
title_run.font.color.rgb = INK

add_para(
    "Two questions came up that couldn't be answered without the system in front "
    "of us. Now it exists \u2014 here is exactly how it works.",
    size=13, color=MUTED, italic=False, space_after=14, line_spacing=1.35,
)

meta_p = doc.add_paragraph()
meta_p.paragraph_format.space_after = Pt(4)
meta_parts = [
    ("Prepared  ", MUTED, False), ("15 September 2026", INK, True),
    ("      Covers  ", MUTED, False), ("Knowledge base & risk model", INK, True),
    ("      Status  ", MUTED, False), ("Implementation complete", INK, True),
]
for text, color, bold in meta_parts:
    r = meta_p.add_run(text)
    r.font.name = MONO_FONT
    r.font.size = Pt(9)
    r.font.color.rgb = color
    r.font.bold = bold

add_hr()
add_para("", size=2, space_after=6)

# ---- at a glance ----
glance_table = doc.add_table(rows=1, cols=3)
glance_items = [
    ("Grounded, not generic", "The chat assistant only answers from a vetted medical library \u2014 never open internet knowledge."),
    ("Transparent scoring", "Every risk result traces back to the specific answers that produced it."),
    ("No invented numbers", "The app never shows a precise-sounding statistic it has no real data to support."),
]
for i, (label, desc) in enumerate(glance_items):
    cell = glance_table.cell(0, i)
    set_cell_shading(cell, "FFFFFF")
    set_cell_borders(cell)
    cell.width = Inches(2.0)
    check_run = cell.paragraphs[0].add_run("\u2713")
    check_run.font.size = Pt(14)
    check_run.font.bold = True
    check_run.font.color.rgb = TEAL
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(label)
    r2.font.bold = True
    r2.font.size = Pt(10)
    r2.font.color.rgb = INK
    r2.font.name = BODY_FONT
    p3 = cell.add_paragraph()
    r3 = p3.add_run(desc)
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = MUTED
    r3.font.name = BODY_FONT

add_para("", size=6, space_after=14)

add_para(
    "Two questions were raised that the previous state of the project couldn't "
    "answer: how does the chatbot's knowledge base actually work, and what were "
    "the design choices and results behind the risk model. Both questions were "
    "fair \u2014 at the time, there was no system to point to. That's now been "
    "built. This brief walks through both, in plain terms, with the technical "
    "specifics attached for the record.",
    size=11.5, space_after=6, line_spacing=1.35,
)

# ============================= QUESTION 1 =============================
add_para("", size=6, space_after=6)
add_hr()
add_eyebrow("Question 1")
add_question_block("How does the chatbot's knowledge base work?")

add_heading2("What it actually knows")
add_para(
    "The assistant draws from a small, deliberately curated library of medical "
    "content \u2014 not the open internet, and not the AI's general training "
    "knowledge. The library covers what cervical cancer is, its symptoms, "
    "prevention, the HPV vaccine, screening schedules, risk factors, and how to "
    "interpret test results \u2014 sourced from WHO, CDC, and American Cancer "
    "Society patient guidance. Keeping it narrow and vetted, rather than broad "
    "and scraped, was the deliberate choice: in a health app, getting less "
    "wrong matters more than covering more ground.",
    line_spacing=1.35,
)

add_heading2("How it finds the right answer")
add_para(
    "When a user asks something, the system doesn't send the question straight "
    "to an AI and hope for the best. It first searches the library for the "
    "passages most relevant to that specific question, then hands only those "
    "passages to the AI and instructs it to answer strictly from them. If "
    "nothing in the library is relevant enough, the assistant says so honestly "
    "rather than guessing.",
    line_spacing=1.35,
)

add_flow([
    "User asks a question, in any of the app's four languages",
    "System searches the vetted library for the closest matches",
    "Only genuinely relevant passages are kept",
    "The AI writes a reply using only those passages, in the user's language",
])

add_heading2("Staying on topic")
add_para(
    "Anything unrelated to cervical cancer or reproductive health \u2014 weather, "
    "politics, entertainment \u2014 is redirected before it ever reaches the AI. "
    "The assistant is scoped on purpose; it isn't meant to be a general chatbot.",
    line_spacing=1.35,
)

add_spec_table(
    ["Component", "What was chosen", "Why"],
    [
        ["Search method", "Local semantic search (embeddings)", "Runs on every message with no per-request cost or external dependency"],
        ["Answer writer", "OpenAI (GPT-4o mini)", "Fluent, natural multilingual replies \u2014 the one place a large hosted model earns its cost"],
        ["Language support", "English & Spanish (full), Swahili & Yoruba (fallback)", "Matches the app's four supported languages; translated content is still expanding"],
        ["Unknown questions", "Declines rather than guesses", "A wrong confident answer is worse than an honest \u201cI don't have that\u201d"],
    ],
)

add_glossary("In plain terms", [
    ("Semantic search", "matching by meaning rather than exact keywords, so \u201cdoes the shot hurt\u201d still finds the vaccine article."),
    ("RAG", "\u201cRetrieval-Augmented Generation\u201d \u2014 the technical name for look-it-up-then-answer, instead of answer-from-memory."),
])

# ============================= QUESTION 2 =============================
add_hr()
add_eyebrow("Question 2")
add_question_block("What are the risk model's design choices \u2014 and what are its results?")

add_heading2("Why it isn't a trained prediction model")
add_para(
    "There is a public research dataset this kind of tool is often built on \u2014 "
    "but it only covers 858 people, with very few confirmed cancer cases and a "
    "lot of missing answers. Training a predictive model on data that thin "
    "tends to produce results that look precise but aren't actually reliable "
    "\u2014 a real risk in a health context. So instead of training a model to "
    "guess, the system was built to reason transparently from established "
    "medical risk factors.",
    line_spacing=1.35,
)

add_heading2("How a score is actually produced")
add_para(
    "Each answer a user gives is checked against a set of known risk factors "
    "\u2014 prior test results, HPV or HIV status, smoking history, screening "
    "history, and others \u2014 and contributes points according to how "
    "strongly that factor is linked to cervical cancer risk in medical "
    "guidance. The points add up to a score, which maps to a Low, Moderate, or "
    "High result.",
    line_spacing=1.35,
)

add_flow([
    "User completes the 30-question assessment",
    "Each answer is checked against known clinical risk factors",
    "Matching factors add weighted points, openly",
    "Result shows the level and exactly which factors drove it",
])

add_heading2("Why there's no \u201c1 in 12\u201d style statistic")
add_para(
    "An early option was showing a precise-sounding figure \u2014 something like "
    "\u201cyou have a 1 in 12 chance.\u201d That was deliberately rejected. The "
    "system has no data connecting its point total to real-world incidence "
    "rates, so a number like that would look far more certain than it "
    "actually is. Instead, the result gives a clear risk level and the "
    "specific reasons behind it \u2014 accurate to what the system actually "
    "knows.",
    line_spacing=1.35,
)

add_heading2("What \u201cresults\u201d means for a system like this")
add_para(
    "Because nothing was trained on data, there's no accuracy percentage to "
    "report the way there would be for a machine-learning model \u2014 and "
    "that's a stated design decision, not a gap. Every completed assessment is "
    "saved, on purpose, so that if the weighting ever needs to be tuned "
    "against real usage patterns, the data to do that already exists.",
    line_spacing=1.35,
)

add_spec_table(
    ["Factor group", "Example", "Weighting logic"],
    [
        ["Test results", "Positive biopsy, abnormal Pap smear", "Heaviest weight \u2014 closest signal to an actual clinical finding"],
        ["Infection history", "HPV, HIV/AIDS status", "Moderate weight \u2014 well-evidenced contributing factors"],
        ["Lifestyle", "Smoking duration, contraceptive use", "Moderate weight, scaled by duration/intensity"],
        ["History", "Partner count, screening history", "Lighter weight \u2014 contributing, not determinative"],
    ],
)

add_glossary("In plain terms", [
    ("Rule-based", "the logic is written and reviewable by a person, not learned automatically from data."),
    ("Calibrated", "a number that's been checked against real outcomes to confirm it means what it claims to mean."),
])

# ============================= STATUS =============================
add_hr()
add_eyebrow("Where this stands today")
add_heading2("Implementation status")

status_items = [
    (TEAL, "Both systems above are fully built and pass all code-level checks."),
    (GOLD, "Final integration testing against a live database is the one step still ahead before this is demo-ready."),
]
for color, text in status_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    dot = p.add_run("\u25CF  ")
    dot.font.color.rgb = color
    dot.font.size = Pt(9)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = INK
    r.font.name = BODY_FONT

add_para("", size=10, space_after=10)
add_hr()
footer_p = doc.add_paragraph()
footer_p.paragraph_format.space_before = Pt(8)
fr = footer_p.add_run("Technical companion to the verbal walkthrough \u2014 cervical cancer awareness app.")
fr.font.name = MONO_FONT
fr.font.size = Pt(8.5)
fr.font.color.rgb = MUTED

out_path = r"C:\Users\Divine\Documents\projects\canproject\Behind-the-Assistant.docx"
doc.save(out_path)
print("Saved:", out_path)
